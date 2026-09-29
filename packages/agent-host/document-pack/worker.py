"""Turnsu documents: fixed operations, no model-supplied Python or shell."""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

LIMIT = 500_000


def office_check(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
            raise ValueError('Office 解压内容超过上限。')
        names = set()
        for entry in entries:
            name = entry.filename.replace('\\', '/')
            if name in names or name.startswith('/') or '..' in name.split('/') or entry.flag_bits & 1:
                raise ValueError('Office 包含不安全或加密条目。')
            names.add(name)
            if any(s in name.lower() for s in ('vbaproject', 'macrosheets', 'externallinks', 'embeddings/')):
                raise ValueError('请提供不含宏、嵌入程序或外部链接的副本。')
            if entry.file_size > max(1, entry.compress_size) * 1000:
                raise ValueError('Office 压缩比例异常。')
            if name.endswith('.rels'):
                data = archive.read(entry)
                if b'TargetMode="External"' in data or b"TargetMode='External'" in data:
                    raise ValueError('文件含有外部资源关系，请先移除后使用文件工具。')


def bounded(value):
    encoded = json.dumps(value, ensure_ascii=False, default=str)
    if len(encoded) > LIMIT:
        raise ValueError('结果超过单次读取上限，请缩小页码、工作表或行范围。')
    return value


def read_document(source, options):
    suffix = source.suffix.lower()
    if suffix in ('.xlsx', '.docx'):
        office_check(source)
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        formula_book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        values_book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            if len(formula_book.sheetnames) > 64:
                raise ValueError('工作表数量超过 64。')
            result = []
            start = int(options.get('startRow', 1))
            count = int(options.get('rows', 100))
            if start < 1 or not 1 <= count <= 1000:
                raise ValueError('每次可读取 1 至 1000 行。')
            selected = options.get('sheet')
            if selected and selected not in formula_book.sheetnames:
                raise ValueError('工作表不存在。')
            for sheet in formula_book:
                if selected and sheet.title != selected:
                    continue
                if sheet.max_column and sheet.max_column > 256:
                    raise ValueError('工作表列数超过 256，请提供较小副本。')
                cached = values_book[sheet.title]
                rows = []
                values = cached.iter_rows(min_row=start, max_row=start + count - 1, max_col=sheet.max_column or 1)
                for row, value_row in zip(sheet.iter_rows(min_row=start, max_row=start + count - 1, max_col=sheet.max_column or 1), values):
                    cells = []
                    for cell, value_cell in zip(row, value_row):
                        if cell.value is not None:
                            cells.append({'cell': cell.coordinate, 'value': value_cell.value if cell.data_type == 'f' else cell.value,
                                          'formula': cell.value if cell.data_type == 'f' else None,
                                          'valueSource': 'saved_cache_not_recalculated' if cell.data_type == 'f' else 'literal', 'numberFormat': cell.number_format})
                    if cells:
                        rows.append(cells)
                result.append({'sheet': sheet.title, 'rows': rows, 'startRow': start, 'endRow': start + count - 1, 'totalRows': sheet.max_row, 'hasMore': (sheet.max_row or 0) >= start + count})
            return bounded({'kind': 'spreadsheet', 'sheets': result, 'emptyCells': 'omitted_with_exact_nonempty_addresses', 'warnings': ['公式值来自文件缓存，未重新计算。']})
        finally:
            formula_book.close()
            values_book.close()
    if suffix == '.docx':
        from docx import Document
        document = Document(source)
        start, count = int(options.get('startParagraph', 1)), int(options.get('paragraphs', 100))
        table_start, row_start, row_count = int(options.get('table', 1)), int(options.get('startRow', 1)), int(options.get('rows', 100))
        if start < 1 or not 1 <= count <= 500 or table_start < 1 or row_start < 1 or not 1 <= row_count <= 500:
            raise ValueError('段落和表格起点从 1 开始，每次最多 500 段/行。')
        all_paragraphs = document.paragraphs
        paragraphs = [{'paragraph': i + 1, 'text': p.text, 'style': p.style.name} for i, p in enumerate(all_paragraphs[start-1:start-1+count], start-1)]
        tables = [{'table': i+1, 'startRow': row_start, 'totalRows': len(table.rows), 'hasMore': row_start-1+row_count<len(table.rows), 'rows': [[cell.text for cell in row.cells] for row in table.rows[row_start-1:row_start-1+row_count]]} for i, table in enumerate(document.tables[table_start-1:table_start+4], table_start-1)]
        return bounded({'kind': 'document', 'paragraphs': paragraphs, 'tables': tables, 'totalParagraphs': len(all_paragraphs), 'totalTables': len(document.tables), 'hasMore': start-1+count<len(all_paragraphs) or table_start+4<len(document.tables),
                        'warnings': ['文本读取不包含图片内容、文本框、批注或精确分页；需要视觉核对时转换并预览 PDF。']})
    if suffix == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(source)
        if pdf.is_encrypted:
            raise ValueError('加密 PDF 需要用户先提供可读取的副本。')
        if len(pdf.pages) > 200:
            raise ValueError('PDF 超过 200 页，请分批处理。')
        start, count = int(options.get('startPage', 1)), int(options.get('pages', 20))
        if start < 1 or not 1 <= count <= 100:
            raise ValueError('每次可读取 1 至 100 页。')
        pages = []
        for i in range(start - 1, min(start - 1 + count, len(pdf.pages))):
            text = pdf.pages[i].extract_text() or ''
            mode = 'text_layer'
            if not text.strip() and options.get('ocr'):
                if not shutil.which('pdftoppm') or not shutil.which('tesseract'):
                    raise ValueError('OCR 组件未安装；请在文件能力中启用 OCR。')
                with tempfile.TemporaryDirectory() as directory:
                    target = str(Path(directory) / 'page')
                    subprocess.run(['pdftoppm', '-f', str(i + 1), '-l', str(i + 1), '-singlefile', '-scale-to', '2400', '-png', str(source), target], check=True, capture_output=True, timeout=30)
                    text = subprocess.run(['tesseract', target + '.png', 'stdout', '-l', 'chi_sim+eng'], check=True, capture_output=True, timeout=60).stdout.decode('utf8')
                    mode = 'ocr_needs_review'
            pages.append({'page': i + 1, 'text': text, 'coverage': mode if text.strip() else 'needs_ocr'})
        return bounded({'kind': 'pdf', 'totalPages': len(pdf.pages), 'pages': pages, 'hasMore': start - 1 + count < len(pdf.pages), 'warnings': ['图片、图表和视觉布局需打开原件核对。']})
    raise ValueError('文件工具支持 .xlsx、.docx 和 .pdf。')


def create_document(target, content):
    suffix = target.suffix.lower()
    if suffix == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        book.remove(book.active)
        sheets = content.get('sheets', [])
        if not 1 <= len(sheets) <= 64:
            raise ValueError('请提供 1 至 64 个工作表。')
        for spec in sheets:
            sheet = book.create_sheet(spec['name'])
            rows = spec.get('rows', [])
            if len(rows) > 10000 or any(not isinstance(row, list) or len(row) > 256 for row in rows):
                raise ValueError('表格数据超过上限。')
            for row in rows:
                for value in row:
                    if isinstance(value, str) and value.startswith('=') and re.search(r'\[|https?:|WEBSERVICE|HYPERLINK|DDE', value, re.I):
                        raise ValueError('不允许生成访问外部资源的公式。')
                sheet.append(row)
            if rows:
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill('solid', fgColor='E8EEE6')
                sheet.freeze_panes = 'A2'
        book.save(target)
        return ['公式已写入，缓存值尚未重算。']
    if suffix == '.docx':
        from docx import Document
        document = Document()
        if content.get('title'):
            document.add_heading(content['title'], 0)
        for paragraph in content.get('paragraphs', []):
            document.add_paragraph(str(paragraph))
        for rows in content.get('tables', []):
            if not rows or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError('新建 Word 表格必须具有一致列数。')
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
        document.save(target)
        return ['请预览核对分页与表格版式。']
    if suffix == '.pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from xml.sax.saxutils import escape
        pdfmetrics.registerFont(TTFont('TurnsuCJK', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', subfontIndex=0))
        style = ParagraphStyle('body', fontName='TurnsuCJK', fontSize=11, leading=18, wordWrap='CJK')
        title_style = ParagraphStyle('title', parent=style, fontSize=18, leading=26, spaceAfter=14)
        story = [Paragraph(escape(str(content['title'])), title_style)] if content.get('title') else []
        for block in content.get('paragraphs', []):
            story.extend([Paragraph(escape(str(block)).replace('\n', '<br/>'), style), Spacer(1, 10)])
        SimpleDocTemplate(str(target), title=str(content.get('title', ''))).build(story)
        return ['PDF 已生成；请打开核对字体与分页。']
    raise ValueError('输出格式必须为 .xlsx、.docx 或 .pdf。')


def edit_document(source, target, changes):
    if source.suffix.lower() != target.suffix.lower():
        raise ValueError('有限编辑必须保持原格式，另存为新文件。')
    suffix = source.suffix.lower()
    if suffix in ('.xlsx', '.docx'):
        office_check(source)
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(source, keep_links=False)
        for change in changes.get('cells', []):
            if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', change['cell']):
                raise ValueError('单元格地址无效。')
            value = change.get('value')
            if isinstance(value, str) and value.startswith('=') and re.search(r'\[|https?:|WEBSERVICE|HYPERLINK|DDE', value, re.I):
                raise ValueError('不允许外部公式。')
            book[change['sheet']][change['cell']] = value
        book.save(target)
        book.close()
        return ['已另存。图形等不受解析库支持的元素可能变化；公式未重算，请核对副本。']
    if suffix == '.docx':
        from docx import Document
        document = Document(source)
        for change in changes.get('paragraphs', []):
            position = int(change['paragraph'])
            if position < 1 or position > len(document.paragraphs):
                raise ValueError('段落位置不在原文件中。')
            document.paragraphs[position - 1].text = str(change['text'])
        document.save(target)
        return ['仅替换指定段落；这些段落的行内格式将重新设置，请核对副本。']
    if suffix == '.pdf':
        from pypdf import PdfReader, PdfWriter
        reader, writer = PdfReader(source), PdfWriter()
        pages = changes.get('keepPages', [])
        if not pages or len(pages) > 200:
            raise ValueError('请选择需要保留的页码。')
        for page in pages:
            if not isinstance(page, int) or page < 1 or page > len(reader.pages):
                raise ValueError('页码无效。')
            writer.add_page(reader.pages[page - 1])
        writer.write(target)
        return ['已生成所选页面的新 PDF；未编辑页面中的文字。']
    raise ValueError('不支持这种有限编辑。')


def main():
    request = json.loads(sys.stdin.read(1_000_001))
    action = request['operation']
    source = Path(request['source']) if request.get('source') else None
    target = Path(request['target']) if request.get('target') else None
    if source and (not source.is_file() or source.is_symlink() or source.stat().st_size > 32 * 1024 * 1024):
        raise ValueError('输入文件不可用或超过 32 MiB。')
    if target and target.exists():
        raise ValueError('输出已存在，不能覆盖。')
    if action == 'read':
        result = read_document(source, request.get('options', {}))
    elif action == 'create':
        result = {'warnings': create_document(target, request.get('content', {}))}
    elif action == 'edit':
        result = {'warnings': edit_document(source, target, request.get('changes', {}))}
    elif action == 'convert':
        if source.suffix in ('.xlsx', '.docx'):
            office_check(source)
        if not shutil.which('libreoffice'):
            raise ValueError('转换/重算组件尚未安装。')
        with tempfile.TemporaryDirectory() as directory:
            profile = Path(directory, 'profile').as_uri()
            subprocess.run(['libreoffice', '-env:UserInstallation=' + profile, '--headless', '--convert-to', target.suffix[1:], '--outdir', directory, str(source)], check=True, capture_output=True, timeout=90)
            converted = Path(directory, source.stem + target.suffix)
            if not converted.is_file():
                raise ValueError('转换没有产生文件；源文件未改动。')
            shutil.copyfile(converted, target)
        result = {'warnings': ['转换/重算完成；复杂公式与版式仍需核对。']}
    else:
        raise ValueError('不支持这个文档操作。')
    if target:
        if not target.is_file() or not 0 < target.stat().st_size <= 32 * 1024 * 1024:
            raise ValueError('输出文件为空或超过上限。')
        result['bytes'] = target.stat().st_size
    print(json.dumps({'ok': True, 'result': bounded(result)}, ensure_ascii=False, default=str))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(json.dumps({'ok': False, 'error': str(error)[:1000]}, ensure_ascii=False))
        sys.exit(1)
